import io
import os
import unittest
import zipfile
from unittest.mock import patch
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, NumberObject, DecodedStreamObject
from worker import Rejected, parse_docx, parse_pdf, render_docx, convert_pdf


def synthetic_pdf(*, encrypted=False, image_only=False):
    """Tiny in-memory parser inputs; no local fonts or pre-generated files required."""
    writer = PdfWriter()
    for page_number in (1, 2):
        page = writer.add_blank_page(width=612, height=792)
        contents = DecodedStreamObject()
        if image_only:
            image = DecodedStreamObject()
            image.set_data(b"\x80\x80\x80")
            image.update({NameObject("/Type"): NameObject("/XObject"),
                          NameObject("/Subtype"): NameObject("/Image"),
                          NameObject("/Width"): NumberObject(1), NameObject("/Height"): NumberObject(1),
                          NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
                          NameObject("/BitsPerComponent"): NumberObject(8)})
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/XObject"): DictionaryObject({NameObject("/Image1"): image})})
            contents.set_data(b"q 100 0 0 100 72 500 cm /Image1 Do Q")
        else:
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                     NameObject("/Subtype"): NameObject("/Type1"),
                                     NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            contents.set_data(f"BT /F1 12 Tf 72 700 Td (Synthetic safety training page {page_number}.) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = contents
    if encrypted:
        writer.encrypt("synthetic-test-only")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def synthetic_snapshot():
    snapshot = {"title": "安全培训测试卷", "durationMinutes": 30, "totalScore": 100,
                "templateVersion": "exam-paper.v1", "items": []}
    for ordinal, kind in enumerate(("SINGLE_CHOICE", "MULTIPLE_CHOICE", "TRUE_FALSE", "SHORT_ANSWER"), 1):
        content = {"type": kind, "stem": "作业前应检查什么？", "analysis": "解析专用哨兵：先检查防护设备。",
                   "sourceRefs": [{"sourceVersionId": "1", "fragmentId": "1", "quote": "依据专用哨兵"}]}
        if kind in ("SINGLE_CHOICE", "MULTIPLE_CHOICE"):
            content.update(options=[{"id": "A", "text": "检查设备"}, {"id": "B", "text": "报告故障"}],
                           correctOptionIds=["A"] if kind == "SINGLE_CHOICE" else ["A", "B"])
        elif kind == "TRUE_FALSE":
            content["answerBoolean"] = True
        else:
            content.update(referenceAnswer="简答答案专用哨兵", rubric=[{"point": "评分专用哨兵", "weight": 100}])
        snapshot["items"].append({"ordinal": ordinal, "score": 25, "content": content})
    return snapshot


class DocumentWorkerTest(unittest.TestCase):
    def test_docx_preserves_paragraph_and_table_locators(self):
        document = Document()
        document.add_paragraph("合成培训资料：先检查设备。")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "表格中的合成资料"
        output = io.BytesIO()
        document.save(output)
        result = parse_docx(output.getvalue())
        self.assertTrue(result["fragments"])
        kinds = {fragment["locator"]["kind"] for fragment in result["fragments"]}
        self.assertIn("DOCX_PARAGRAPH", kinds)
        self.assertIn("DOCX_TABLE_CELL", kinds)

    def test_duplicate_zip_is_rejected_before_xml_parser(self):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            archive.writestr("word/document.xml", "first")
            archive.writestr("word/document.xml", "second")
        with self.assertRaises(Rejected):
            parse_docx(output.getvalue())

    def test_text_boxes_and_deleted_text_are_not_silently_included(self):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>可用正文</w:t></w:r><w:txbxContent><w:p><w:r><w:t>文本框秘密</w:t></w:r></w:p></w:txbxContent><w:del><w:r><w:t>已删除</w:t></w:r></w:del></w:p></w:body></w:document>')
        result = parse_docx(output.getvalue())
        self.assertEqual("可用正文", result["fragments"][0]["text"])

    def test_external_relationships_and_dtd_are_rejected(self):
        for document, relations in [(b'<!DOCTYPE x [<!ENTITY xxe SYSTEM "file:///etc/passwd">]><x>&xxe;</x>', b'<r/>'),
                                    (b'<x/>', b'<Relationships><Relationship TargetMode="External" Target="https://example.invalid"/></Relationships>')]:
            output = io.BytesIO()
            with zipfile.ZipFile(output, "w") as archive:
                archive.writestr("word/document.xml", document)
                archive.writestr("[Content_Types].xml", b'<Types/>')
                archive.writestr("word/_rels/document.xml.rels", relations)
            with self.assertRaises(Rejected):
                parse_docx(output.getvalue())

    def test_pdf_pages_have_exact_page_locators(self):
        result = parse_pdf(synthetic_pdf())
        self.assertEqual({1, 2}, {part["locator"]["page"] for part in result["fragments"]})
        self.assertTrue(all("Synthetic safety training" in part["text"] for part in result["fragments"]))

    def test_encrypted_and_scan_only_are_rejected(self):
        for mode in ("encrypted", "image_only"):
            with self.subTest(mode=mode), self.assertRaises(Rejected):
                parse_pdf(synthetic_pdf(**{mode: True}))

    def test_export_snapshot_produces_clean_student_and_teacher_docx(self):
        snapshot = synthetic_snapshot()
        for audience in ("STUDENT", "TEACHER"):
            data = render_docx(snapshot, audience)
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")
                self.assertIn("w:numPr", xml)
                from lxml import etree
                styles = etree.fromstring(archive.read("word/styles.xml"))
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                title = styles.xpath("//w:style[@w:styleId='Title']", namespaces=ns)[0]
                self.assertFalse(title.xpath("./w:pPr/w:pBdr", namespaces=ns))
                self.assertFalse(title.xpath("./w:rPr/w:rFonts/@w:eastAsiaTheme", namespaces=ns))
                if audience == "STUDENT":
                    self.assertNotIn("参考答案", xml)
                    self.assertNotIn("评分点", xml)
                    self.assertNotIn("资料版本", xml)
                    for sentinel in ("解析专用哨兵", "依据专用哨兵", "简答答案专用哨兵", "评分专用哨兵"):
                        self.assertNotIn(sentinel, xml)
                    body = etree.fromstring(archive.read("word/document.xml"))
                    lines = body.xpath("//w:p[w:r/w:t='________________________________________________________________']", namespaces=ns)
                    self.assertEqual(4, len(lines))
                    self.assertTrue(all(line.xpath("./w:pPr/w:keepNext", namespaces=ns) for line in lines[:3]))
                else:
                    self.assertIn("参考答案", xml)
                    self.assertIn("评分点", xml)
                    for sentinel in ("解析专用哨兵", "依据专用哨兵", "简答答案专用哨兵", "评分专用哨兵"):
                        self.assertIn(sentinel, xml)

    def test_conversion_missing_is_explicit_not_fake_pdf(self):
        with patch.dict(os.environ, {}, clear=True), patch("worker.shutil.which", return_value=None):
            with self.assertRaises(Rejected) as error:
                convert_pdf(b"not converted")
        self.assertEqual("EXAM_CONVERTER_MISSING", error.exception.code)


if __name__ == "__main__":
    unittest.main()
